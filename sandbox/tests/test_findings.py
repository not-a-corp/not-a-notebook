"""The README's data-quality rules, one at a time, through the whole profiler.

Each test writes a small file with one known problem, profiles it the way the
API does, and checks the finding: its kind, where it points, and that suspicious
values are shown both ways. One test asks the opposite — that a clean file is
met with silence, because a profiler that cries wolf trains the model to ignore
it.

These files were written to exercise the detectors, so they prove the detectors
fire. How well the profile serves real dirty files is measured separately, on
files nobody here wrote (Phase 9).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import openpyxl
import pytest
from not_a_notebook_findings import date_reading, readings
from not_a_notebook_profile import profile_of


def write(tmp_path: Path, name: str, text: str, encoding: str = "utf-8") -> str:
    path = tmp_path / name
    path.write_bytes(text.encode(encoding))

    return str(path)


def findings(profile: dict[str, Any], kind: str) -> list[dict[str, Any]]:
    found = []

    for item in profile["findings"]:
        if item["kind"] == kind:
            found.append(item)

    for table in profile["tables"]:
        for item in table["findings"]:
            if item["kind"] == kind:
                found.append(item)

    return found


def every_kind(profile: dict[str, Any]) -> list[str]:
    kinds = []

    for item in profile["findings"]:
        kinds.append(item["kind"])

    for table in profile["tables"]:
        for item in table["findings"]:
            kinds.append(item["kind"])

    return kinds


# ── a clean file says nothing ────────────────────────────────────────────────


def test_a_clean_file_has_no_findings(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "clean.csv",
        "region,quantity,price,revenue,date\n"
        "North,2,10.5,21.0,2025-01-03\n"
        "South,1,99.0,99.0,2025-01-04\n"
        "East,4,5.25,21.0,2025-01-05\n"
        "West,3,7.0,21.0,2025-01-06\n"
        "Centre,5,2.0,10.0,2025-01-07\n",
    )

    profile = profile_of(path)

    assert every_kind(profile) == []


# ── subtotal and grand-total rows ────────────────────────────────────────────


def test_a_total_row_among_the_records_is_found_and_checked(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "vendas.csv",
        "região;vendas\nSudeste;1.840.000,00\nNordeste;512.000,50\nTOTAL;2.352.000,50\n",
        encoding="latin-1",
    )

    found = findings(profile_of(path), "subtotal_rows")

    assert len(found) == 1
    # Line 4 of the file: header on 1, records on 2 and 3.
    assert found[0]["rows"] == [4]
    assert found[0]["sums_checked"] == [{"row": 4, "column": "vendas"}]


def test_subtotals_between_groups_are_each_found(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "groups.csv",
        "store,item,sales\n"
        "A,x,10\nA,y,5\nSubtotal A,,15\n"
        "B,x,7\nB,y,3\nSubtotal B,,10\n"
        "Total geral,,25\n",
    )

    found = findings(profile_of(path), "subtotal_rows")[0]

    assert found["rows"] == [4, 7, 8]
    assert {"row": 4, "column": "sales"} in found["sums_checked"]
    assert {"row": 7, "column": "sales"} in found["sums_checked"]


# ── a header that is not on the first row ────────────────────────────────────


def test_a_title_block_above_the_header_is_found_and_skipped(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "report.csv",
        "Relatório de vendas\nGerado em 03/10/2026\n\nregião;vendas\nSudeste;10\nNordeste;7\n",
    )

    profile = profile_of(path)
    found = findings(profile, "header_not_on_first_row")
    table = profile["tables"][0]

    assert found[0]["header_row"] == 4
    assert found[0]["above"] == ["Relatório de vendas", "Gerado em 03/10/2026"]
    assert [column["name"] for column in table["columns"]] == ["região", "vendas"]
    assert table["rows"] == 2


def test_a_header_lower_down_in_a_sheet_is_found(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Vendas por região — 2025"])
    sheet.append([])
    sheet.append(["região", "vendas"])
    sheet.append(["Sudeste", 10])
    sheet.append(["Nordeste", 7])
    path = tmp_path / "report.xlsx"
    workbook.save(path)

    profile = profile_of(str(path))

    assert findings(profile, "header_not_on_first_row")[0]["header_row"] == 3
    assert profile["tables"][0]["rows"] == 2


# ── the relevant sheet is not the first ──────────────────────────────────────


def test_data_on_a_later_sheet_is_pointed_at(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    notes = workbook.active
    assert notes is not None
    notes.title = "Leia-me"
    notes.append(["Fonte: ERP, exportado à mão"])

    data = workbook.create_sheet("Vendas")
    data.append(["região", "mês", "vendas"])
    for month in range(1, 13):
        data.append(["Sudeste", month, 100 + month])
        data.append(["Nordeste", month, 50 + month])

    path = tmp_path / "vendas.xlsx"
    workbook.save(path)

    found = findings(profile_of(str(path)), "data_not_on_first_sheet")

    assert found[0]["first"] == "Leia-me"
    assert found[0]["largest"] == "Vendas"


# ── exact duplicates ─────────────────────────────────────────────────────────


def test_exact_duplicates_are_counted_and_pointed_at(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "dupes.csv",
        "order,region,sales\n1,N,10\n2,S,20\n1,N,10\n3,N,30\n1,N,10\n",
    )

    found = findings(profile_of(path), "exact_duplicates")[0]

    assert found["count"] == 2
    assert found["rows"] == [4, 6]


# ── mixed number formats ─────────────────────────────────────────────────────


def test_numbers_in_three_formats_are_named(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "mixed.csv",
        "item;price\na;1234.56\nb;1.234,56\nc;R$ 1.234,56\nd;99,90\n",
    )

    found = findings(profile_of(path), "number_formats")[0]

    assert found["column"] == "price"
    assert found["mixed"] is True
    assert found["formats"]["plain"] == 1
    assert found["formats"]["brazilian"] == 3
    assert found["formats"]["with_currency"] == 1


def test_a_value_that_reads_two_ways_is_shown_both_ways(tmp_path: Path) -> None:
    path = write(tmp_path, "ambiguous.csv", "item;weight\na;1.234\nb;2.500,75\nc;3,25\n")

    found = findings(profile_of(path), "number_formats")[0]

    assert found["ambiguous"] == [
        {"value": "1.234", "readings": {"plain": 1.234, "brazilian": 1234.0}}
    ]


def test_the_number_readings_themselves() -> None:
    assert readings("1234.56") == {"plain": 1234.56}
    assert readings("1.234,56") == {"brazilian": 1234.56}
    assert readings("R$ 1.234,56") == {"brazilian": 1234.56}
    assert readings("1,234.56") == {"us": 1234.56}
    assert readings("1,234") == {"brazilian": 1.234, "us": 1234.0}
    assert readings("Sudeste") == {}


# ── mixed date formats ───────────────────────────────────────────────────────


def test_day_first_dates_are_recognised_from_one_telling_value(tmp_path: Path) -> None:
    path = write(tmp_path, "dates.csv", "when;sales\n01/02/2025;1\n15/03/2025;2\n")

    found = findings(profile_of(path), "date_formats")[0]

    assert found["formats"] == {"ambiguous": 1, "dd/mm": 1}
    assert found["ambiguous"] == []


def test_dates_that_could_be_either_are_shown_both_ways(tmp_path: Path) -> None:
    path = write(tmp_path, "dates.csv", "when;sales\n01/02/2025;1\n03/04/2025;2\n")

    found = findings(profile_of(path), "date_formats")[0]

    assert found["ambiguous"][0] == {
        "value": "01/02/2025",
        "as_dd_mm": "day 01, month 02",
        "as_mm_dd": "month 01, day 02",
    }


def test_iso_and_slashed_dates_in_one_column_are_mixed(tmp_path: Path) -> None:
    path = write(tmp_path, "dates.csv", "when;sales\n2025-01-03;1\n15/03/2025;2\n2025-01-05;3\n")

    found = findings(profile_of(path), "date_formats")[0]

    assert "The formats are mixed." in found["message"]


def test_the_date_readings_themselves() -> None:
    assert date_reading("2025-03-15") == "iso"
    assert date_reading("15/03/2025") == "dd/mm"
    assert date_reading("03/15/2025") == "mm/dd"
    assert date_reading("03/04/2025") == "ambiguous"
    assert date_reading("Sudeste") is None


# ── inconsistent category labels ─────────────────────────────────────────────


def test_one_label_spelled_several_ways_is_grouped(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "labels.csv",
        "region,sales\nSudeste,1\nsudeste,2\nSudeste ,3\nSUDESTE,4\nNordeste,5\nNordeste,6\n",
    )

    found = findings(profile_of(path), "inconsistent_labels")[0]

    # The trailing space is its own spelling, and is shown as such.
    assert found["column"] == "region"
    assert found["variants"] == [{"Sudeste": 1, "sudeste": 1, "Sudeste ": 1, "SUDESTE": 1}]


def test_accents_do_not_make_a_label_new(tmp_path: Path) -> None:
    path = write(tmp_path, "labels.csv", "city,sales\nSão Paulo,1\nSao Paulo,2\nRecife,3\n")

    found = findings(profile_of(path), "inconsistent_labels")[0]

    assert found["variants"] == [{"São Paulo": 1, "Sao Paulo": 1}]


# ── columns that contradict each other ───────────────────────────────────────


def test_revenue_that_is_not_quantity_times_price_is_found(tmp_path: Path) -> None:
    path = write(
        tmp_path,
        "orders.csv",
        "faturamento,qtd,preço\n20,2,10\n30,3,10\n45,5,9\n99,4,8\n12,3,4\n",
    )

    found = findings(profile_of(path), "contradicting_columns")

    assert len(found) == 1
    assert found[0]["columns"] == ["faturamento", "qtd", "preço"]
    assert found[0]["rows"] == [5]


# ── outliers that dominate totals ────────────────────────────────────────────


def test_one_value_that_is_most_of_the_total_is_found(tmp_path: Path) -> None:
    rows = "".join(f"r{n},{10 + n}\n" for n in range(10))
    path = write(tmp_path, "sales.csv", "store,sales\n" + rows + "typo,1000000\n")

    found = findings(profile_of(path), "dominating_value")[0]

    assert found["column"] == "sales"
    assert found["rows"] == [12]
    assert found["share"] > 0.99


def test_a_total_row_is_not_also_an_outlier(tmp_path: Path) -> None:
    rows = "".join(f"r{n},{10 + n}\n" for n in range(10))
    path = write(tmp_path, "sales.csv", "store,sales\n" + rows + "Total,145\n")

    profile = profile_of(path)

    assert findings(profile, "subtotal_rows")
    assert findings(profile, "dominating_value") == []


# ── missing values ───────────────────────────────────────────────────────────


def test_missing_values_written_as_text_are_found(tmp_path: Path) -> None:
    path = write(tmp_path, "gaps.csv", "region,sales\nN,10\nS,-\nNE,n/d\nCO,7\nSE,12\n")

    found = findings(profile_of(path), "missing_values")[0]

    assert found["column"] == "sales"
    assert found["count"] == 2
    assert found["placeholders"] == {"-": 1, "n/d": 1}


def test_a_mostly_empty_column_is_found(tmp_path: Path) -> None:
    path = write(tmp_path, "gaps.csv", "region,notes\nN,\nS,\nNE,ok\nCO,\n")

    found = findings(profile_of(path), "missing_values")[0]

    assert found["column"] == "notes"
    assert found["count"] == 3


@pytest.mark.parametrize("blank", ["\n", "\n\n"])
def test_rows_keep_their_file_numbers_across_blank_lines(tmp_path: Path, blank: str) -> None:
    path = write(tmp_path, "gaps.csv", f"order,sales\n1,10\n{blank}1,10\n")

    found = findings(profile_of(path), "exact_duplicates")[0]

    assert found["rows"] == [3 + blank.count("\n")]
